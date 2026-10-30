"""ClearBrief — a private, local-first document summarizer."""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
import reportlab
import streamlit as st
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from summarizer import SUPPORTED_EXTENSIONS, analyze_text, extract_text, summarize_text


st.set_page_config(
    page_title="AI-Text Summarization",
    page_icon=":material/auto_awesome:",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
        :root {
            --bg-1: #f7f4ff;
            --bg-2: #efe9ff;
            --bg-3: #e7defd;
            --bg-4: #f3edff;
            --panel: rgba(255, 255, 255, 0.42);
            --panel-strong: rgba(255, 255, 255, 0.66);
            --panel-soft: rgba(255, 255, 255, 0.32);
            --border: rgba(156, 138, 210, 0.24);
            --text-strong: #291f42;
            --text-soft: #4d4b68;
            --text-muted: #696282;
            --primary: #a88dff;
            --primary-strong: #7a5ae6;
            --secondary: #d5c4ff;
            --tertiary: #f0e7ff;
            --shadow: rgba(103, 78, 156, 0.12);
        }

        html, body, .stApp {
            background:
                radial-gradient(circle at 12% 10%, rgba(168, 141, 255, 0.18), transparent 15%),
                radial-gradient(circle at 82% 10%, rgba(122, 90, 230, 0.16), transparent 18%),
                radial-gradient(circle at 70% 80%, rgba(196, 172, 255, 0.18), transparent 20%),
                linear-gradient(180deg, var(--bg-1) 0%, var(--bg-2) 45%, var(--bg-3) 82%, var(--bg-4) 100%);
            color: var(--text-strong);
        }

        .stApp {
            background-attachment: fixed;
        }

        .block-container {
            padding-top: 2.2rem;
            padding-bottom: 3rem;
            padding-left: 1.5rem;
            padding-right: 1.5rem;
            max-width: 1500px;
        }

        [data-testid="stSidebar"] {
            background: rgba(255, 255, 255, 0.34);
            border-right: 1px solid var(--border);
            backdrop-filter: blur(18px);
            -webkit-backdrop-filter: blur(18px);
            box-shadow: 12px 0 30px rgba(120, 99, 162, 0.08);
        }

        [data-testid="stSidebar"] .block-container {
            padding-top: 1.5rem;
            padding-left: 1rem;
            padding-right: 1rem;
        }

        h1, h2, h3, h4, h5, h6 {
            color: var(--text-strong);
            letter-spacing: -0.04em;
            font-weight: 700;
        }

        h1 {
            background: linear-gradient(90deg, #2d2448 0%, #54477b 32%, #7d5fe2 68%, #b49cff 100%);
            -webkit-background-clip: text;
            background-clip: text;
            color: transparent;
        }

        p, li, span, div, label {
            color: var(--text-soft);
        }

        .stTabs [role="tablist"] {
            gap: 0.5rem;
            border-bottom: 1px solid rgba(148, 163, 184, 0.18);
            padding-bottom: 0.35rem;
        }

        .stTabs [role="tab"] {
            background: rgba(255, 255, 255, 0.22);
            color: var(--text-muted);
            border: 1px solid rgba(148, 163, 184, 0.18);
            border-radius: 0.9rem 0.9rem 0 0;
            padding: 0.65rem 1rem;
            min-height: 2.8rem;
            backdrop-filter: blur(8px);
        }

        .stTabs [role="tab"][aria-selected="true"] {
            background: linear-gradient(135deg, rgba(168, 141, 255, 0.10), rgba(214, 198, 255, 0.08));
            border-color: rgba(168, 141, 255, 0.18);
            color: var(--text-strong);
            box-shadow: inset 0 1px 0 rgba(255,255,255,0.5);
        }

        .stButton > button,
        .stDownloadButton > button {
            border: 1px solid rgba(122, 90, 230, 0.18);
            border-radius: 0.9rem;
            background: linear-gradient(135deg, #9d84ff 0%, #7e62ea 52%, #684ad1 100%);
            color: white;
            font-weight: 700;
            letter-spacing: 0.01em;
            box-shadow: 0 12px 28px rgba(122, 90, 230, 0.16);
            transition: transform 0.18s ease, box-shadow 0.18s ease, filter 0.18s ease;
        }

        .stButton > button:hover,
        .stDownloadButton > button:hover {
            transform: translateY(-1px);
            box-shadow: 0 16px 30px rgba(122, 90, 230, 0.22);
            filter: brightness(1.02);
        }

        .stButton > button:focus,
        .stDownloadButton > button:focus {
            box-shadow: 0 0 0 0.2rem rgba(122,90,230,0.18);
        }

        .stTextInput > div > div > input,
        .stTextArea textarea,
        .stSelectbox > div > div,
        .stNumberInput input,
        .stDateInput input,
        .stFileUploader > section,
        .stSelectbox select,
        .stMultiSelect > div > div {
            background: rgba(255, 255, 255, 0.24);
            color: var(--text-strong);
            border: 1px solid rgba(148, 163, 184, 0.24);
            border-radius: 0.9rem;
            box-shadow: inset 0 1px 1px rgba(255,255,255,0.18);
            backdrop-filter: blur(10px);
        }

        .stTextInput label,
        .stTextArea label,
        .stSelectbox label,
        .stNumberInput label,
        .stFileUploader label {
            color: var(--text-soft);
            font-weight: 600;
        }

        [data-testid="stMetric"] {
            background: rgba(255, 255, 255, 0.30);
            border: 1px solid rgba(255, 255, 255, 0.22);
            border-radius: 1rem;
            box-shadow: 0 12px 26px rgba(148, 163, 184, 0.12);
            padding: 1rem 1.05rem;
            backdrop-filter: blur(16px);
        }

        [data-testid="stMetricLabel"] {
            color: var(--text-muted);
            font-size: 0.72rem;
            letter-spacing: 0.08em;
            text-transform: uppercase;
        }

        [data-testid="stMetricValue"] {
            color: var(--text-strong);
            font-weight: 700;
            font-size: 1.7rem;
        }

        [data-testid="stVerticalBlockBorderWrapper"],
        [data-testid="stContainer"],
        [data-testid="stExpander"],
        .stAlert,
        .stDataFrame {
            background: rgba(255, 255, 255, 0.18);
            border: 1px solid rgba(255, 255, 255, 0.24);
            border-radius: 1rem;
            box-shadow: 0 12px 26px rgba(148, 163, 184, 0.10);
            backdrop-filter: blur(16px);
        }

        [data-testid="stDataFrame"] {
            border-radius: 0.9rem;
            overflow: hidden;
        }

        .stAlert {
            padding: 0.9rem 1rem;
            border-radius: 0.9rem;
        }

        .stProgress > div > div > div {
            background: linear-gradient(90deg, var(--primary), var(--primary-strong), var(--secondary));
            border-radius: 999px;
        }

        div[data-testid="stDecoration"] {
            background: linear-gradient(90deg, rgba(168, 141, 255, 0.42), rgba(122, 90, 230, 0.32), rgba(213, 196, 255, 0.24));
            border-radius: 999px;
        }

        .stMarkdown a {
            color: #5a47d2;
        }

        .stMarkdown code {
            background: rgba(255, 255, 255, 0.22);
            color: #334155;
            border: 1px solid rgba(148, 163, 184, 0.18);
            border-radius: 0.45rem;
            padding: 0.15rem 0.35rem;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

for key, default in {
    "source_text": "",
    "current_result": None,
    "result_source_hash": None,
    "saved_documents": [],
    "upload_cache": {},
    "active_upload_key": None,
    "summary_options_key": None,
}.items():
    if key not in st.session_state:
        st.session_state[key] = default


def _fingerprint(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _make_docx(title: str, summary: str) -> bytes:
    document = Document()
    document.add_heading(title, 0)
    for paragraph in summary.splitlines():
        if paragraph.strip():
            document.add_paragraph(paragraph, style="List Bullet" if paragraph.startswith(("•", "☐")) else "Normal")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _make_pdf(title: str, summary: str) -> bytes:
    font_dir = Path(reportlab.__file__).parent / "fonts"
    pdfmetrics.registerFont(TTFont("Vera", str(font_dir / "Vera.ttf")))
    pdfmetrics.registerFont(TTFont("Vera-Bold", str(font_dir / "VeraBd.ttf")))
    styles = getSampleStyleSheet()
    styles["Title"].fontName = "Vera-Bold"
    styles["BodyText"].fontName = "Vera"

    content = [
        Paragraph(escape(title), styles["Title"]),
        Spacer(1, 12),
    ]
    for line in summary.splitlines():
        if line.strip():
            text = line.strip()
            if text.startswith("☐"):
                text = "[ ]" + text[1:]
            elif text.startswith("•"):
                text = "-" + text[1:]
            content.append(Paragraph(escape(text), styles["BodyText"]))
            content.append(Spacer(1, 6))

    stream = io.BytesIO()
    SimpleDocTemplate(stream, pagesize=letter).build(content)
    return stream.getvalue()


with st.sidebar:
    st.title("✦ AI-Text Summarization")
    st.caption("Document intelligence")
    st.divider()
    st.caption("Summary settings")
    summary_style = st.selectbox(
        "Output format",
        ["Balanced", "Bullet points", "Executive brief", "Action items"],
        help="Choose how the extracted summary is presented.",
    )
    summary_length = st.segmented_control(
        "Summary length", ["Brief", "Standard", "Detailed"], default="Standard"
    )
    focus = st.text_input(
        "Focus on a topic (optional)",
        placeholder="e.g. risks, recommendations, Q3 results",
        help="Give relevant sentences a ranking boost. The summary still uses only the source text.",
    )
    st.divider()
    st.caption("Add documents")
    uploaded_files = st.file_uploader(
        "Upload one or more files",
        type=[ext.lstrip(".") for ext in sorted(SUPPORTED_EXTENSIONS)],
        accept_multiple_files=True,
        help="Text, PDF, Word, PowerPoint, Excel, web, data, RTF, and image files (OCR). 50 MB per file.",
    )
    extracted_docs: dict[str, str] = {}
    upload_errors: list[str] = []
    for uploaded in uploaded_files or []:
        if uploaded.size > 50 * 1024 * 1024:
            upload_errors.append(f"{uploaded.name}: exceeds the 50 MB per-file limit.")
            continue
        payload = uploaded.getvalue()
        digest = hashlib.sha256(payload).hexdigest()
        cache_key = f"{uploaded.name}:{digest}"
        if cache_key in st.session_state.upload_cache:
            text = st.session_state.upload_cache[cache_key]
        else:
            try:
                text = extract_text(uploaded.name, payload)
                st.session_state.upload_cache[cache_key] = text
            except Exception as exc:
                upload_errors.append(f"{uploaded.name}: {exc}")
                continue
        if not text.strip():
            upload_errors.append(f"{uploaded.name}: no readable text could be extracted.")
            continue
        extracted_docs[uploaded.name] = text
    for message in upload_errors:
        st.warning(message)

    if extracted_docs:
        document_names = list(extracted_docs)
        chosen_document = st.selectbox("Choose a document", document_names)
        chosen_text = extracted_docs[chosen_document]
        chosen_hash = _fingerprint(chosen_text)
        upload_key = f"{chosen_document}:{chosen_hash}"
        summary_options = (summary_style, summary_length or "Standard", focus)
        if st.session_state.active_upload_key != upload_key:
            st.session_state.source_text = chosen_text
            st.session_state.active_source_name = chosen_document
            st.session_state.active_upload_key = upload_key
            st.session_state.current_result = summarize_text(
                chosen_text, style=summary_style, length=summary_length or "Standard", focus=focus
            )
            st.session_state.result_source_hash = chosen_hash
            st.session_state.result_source_name = chosen_document
            st.session_state.summary_options_key = summary_options
        elif st.session_state.source_text == chosen_text and (
            st.session_state.result_source_hash != chosen_hash
            or st.session_state.summary_options_key != summary_options
        ):
            st.session_state.current_result = summarize_text(
                chosen_text, style=summary_style, length=summary_length or "Standard", focus=focus
            )
            st.session_state.result_source_hash = chosen_hash
            st.session_state.result_source_name = chosen_document
            st.session_state.summary_options_key = summary_options
        st.caption(f"{len(document_names)} file(s) ready · selected file summarized automatically")
        with st.expander("Batch summarize all uploaded files"):
            batch_length = st.selectbox("Batch length", ["Brief", "Standard", "Detailed"], key="batch_length")
            if st.button("Create batch summaries", icon=":material/library_books:", width="stretch"):
                batch = []
                for name, text in extracted_docs.items():
                    result = summarize_text(text, style=summary_style, length=batch_length, focus=focus)
                    batch.append({"Document": name, "Words": analyze_text(text)["words"], "Summary": result["summary"]})
                st.session_state.batch_results = batch
                st.success(f"Created summaries for {len(batch)} document(s).")

    st.info("Local-first processing. Documents are not sent to an AI service.", icon=":material/lock:")


st.caption("YOUR IDEAS, CLARIFIED")
st.title("AI-Text Summarization")
st.write("Turn long-form content into clear takeaways. Paste text or upload documents, then explore the details behind every summary.")

source_col, result_col = st.columns([1.04, 0.96], gap="large")
with source_col:
    st.caption("01 / SOURCE MATERIAL")
    sample_dataset_path = Path(__file__).parent / "sample_data" / "text_summarization_samples.json"
    sample_documents = json.loads(sample_dataset_path.read_text(encoding="utf-8"))
    sample_titles = [document["title"] for document in sample_documents]
    selected_sample_title = st.selectbox("Sample text", sample_titles)
    if st.button("Load sample text", icon=":material/menu_book:", width="content"):
        selected_sample = next(
            document for document in sample_documents if document["title"] == selected_sample_title
        )
        st.session_state.source_text = selected_sample["text"]
        st.session_state.active_source_name = selected_sample_title
        st.session_state.active_upload_key = None
        st.session_state.current_result = None
        st.session_state.result_source_hash = None
    st.text_area(
        "Source text",
        key="source_text",
        height=360,
        placeholder="Paste an article, report, meeting notes, research paper, or any text here…",
        help="You can paste text directly or load a supported document using the sidebar.",
    )
    source_text = st.session_state.source_text
    source_stats = analyze_text(source_text)
    metric_cols = st.columns(4)
    metric_cols[0].metric("Words", f"{source_stats['words']:,}")
    metric_cols[1].metric("Characters", f"{source_stats['characters']:,}")
    metric_cols[2].metric("Sentences", f"{source_stats['sentences']:,}")
    if st.button(
        "Generate summary",
        type="primary",
        icon=":material/auto_awesome:",
        width="stretch",
        disabled=not source_text.strip(),
    ):
        st.session_state.current_result = summarize_text(
            source_text, style=summary_style, length=summary_length or "Standard", focus=focus
        )
        st.session_state.result_source_hash = _fingerprint(source_text)
        st.session_state.result_source_name = getattr(st.session_state, "active_source_name", "Pasted text")
        st.session_state.summary_options_key = (
            summary_style, summary_length or "Standard", focus
        )

with result_col:
    st.caption("02 / YOUR SUMMARY")
    result = st.session_state.current_result
    is_current_result = result and st.session_state.result_source_hash == _fingerprint(source_text)
    if is_current_result and result.get("summary"):
        st.subheader("The essentials, at a glance")
        with st.container(border=True):
            st.text(result["summary"])
        output_stats = analyze_text(result["summary"])
        st.markdown("")
        result_cols = st.columns(3)
        result_cols[0].metric("Key points", len(result["key_points"]))
        result_cols[1].metric("Shorter by", f"{result.get('compression_percent', 0)}%")
        result_cols[2].metric("Summary words", f"{output_stats['words']:,}")
        source_name = getattr(st.session_state, "result_source_name", "Pasted text")
        document_title = Path(source_name).stem or "Pasted text"
        export_basename = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", document_title).strip(" .") or "Pasted text"
        export_cols = st.columns(3)
        export_cols[0].download_button("Plain text", f"{document_title}\n\n{result['summary']}", file_name=f"{export_basename}.txt", mime="text/plain", icon=":material/description:", width="stretch")
        export_cols[1].download_button("Word document", _make_docx(document_title, result["summary"]), file_name=f"{export_basename}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", icon=":material/article:", width="stretch")
        export_cols[2].download_button("PDF", _make_pdf(document_title, result["summary"]), file_name=f"{export_basename}.pdf", mime="application/pdf", icon=":material/picture_as_pdf:", width="stretch")
        save_col = st.container()
        if save_col.button("Save to library", icon=":material/bookmark_add:", width="stretch"):
            st.session_state.saved_documents.append({
                "name": source_name,
                "summary": result["summary"],
                "words": source_stats["words"],
                "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
            })
            st.toast("Saved to your session library.")
    elif source_text.strip() and result:
        st.info("The source has changed. Generate a new summary to refresh this result.", icon=":material/info:")
    else:
        with st.container(border=True):
            st.subheader("A clear view is one click away")
            st.markdown("Your summary will appear here, with key takeaways and ready-to-share exports.")
            st.caption("Choose a format, set a reading length, then generate your summary.")


overview_tab, insights_tab, library_tab = st.tabs(["Overview", "Text insights", "Library"])
with overview_tab:
    st.markdown("#### Key takeaways")
    if is_current_result and result.get("key_points"):
        for point in result["key_points"]:
            st.markdown(f"- {point}")
    else:
        st.caption("Generate a summary to see its key takeaways here.")
    if getattr(st.session_state, "batch_results", None):
        st.markdown("#### Batch results")
        st.dataframe(pd.DataFrame(st.session_state.batch_results), hide_index=True, width="stretch")

with insights_tab:
    st.markdown("#### Document profile")
    if source_text.strip():
        insight_cols = st.columns(4)
        insight_cols[0].metric("Paragraphs", f"{source_stats['paragraphs']:,}")
        insight_cols[1].metric("Avg. sentence", f"{source_stats['average_sentence_length']} words")
        insight_cols[2].metric("Avg. word", f"{source_stats['average_word_length']} chars")
        insight_cols[3].metric("Speaking time", f"{source_stats['speaking_minutes']} min")
        readability = source_stats["flesch_score"]
        readability_label = "Easy" if readability >= 60 else "Moderate" if readability >= 30 else "Advanced"
        st.progress(readability / 100, text=f"Flesch reading ease: {readability:.1f} · {readability_label}")
        st.caption("Readability is a rough English-language estimate; scores are less reliable for non-English text, tables, and fragments.")
        if source_stats["keywords"]:
            keyword_frame = pd.DataFrame(source_stats["keywords"], columns=["Keyword", "Occurrences"])
            keyword_left, keyword_right = st.columns([1, 1.4])
            with keyword_left:
                st.markdown("##### Frequent terms")
                st.dataframe(keyword_frame, hide_index=True, height=320, width="stretch")
            with keyword_right:
                st.markdown("##### Word frequency")
                st.bar_chart(keyword_frame.set_index("Keyword"), horizontal=True, height=320, alt="Most frequent content words in the source text")
    else:
        st.info("Add source material to unlock word counts, reading estimates, readability, and keyword analysis.")

with library_tab:
    st.markdown("#### Saved this session")
    saved = st.session_state.saved_documents
    if saved:
        for index, item in enumerate(reversed(saved)):
            with st.container(border=True):
                st.markdown(f"**{item['name']}** · {item['words']:,} words · saved {item['saved_at']}")
                st.write(item["summary"])
        if st.button("Clear session library", icon=":material/delete_sweep:"):
            st.session_state.saved_documents = []
            st.rerun()
    else:
        st.caption("Summaries you save will be available here until this app session ends.")

st.divider()
st.caption("ClearBrief · Extractive summaries retain source wording. Always verify important details against the original document.")