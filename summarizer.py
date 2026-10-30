"""Local document extraction, analysis, and extractive summarization helpers."""

from __future__ import annotations

import csv
import html
import io
import json
import math
import re
from collections import Counter
from dataclasses import asdict, dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree import ElementTree


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json",
    ".yaml", ".yml", ".xml", ".html", ".htm", ".rtf", ".pdf", ".docx",
    ".pptx", ".xlsx", ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif",
    ".tiff",
}

STOP_WORDS = set(
    "a an the and or but if while of at by for with about against between into through during before after above below to from up down in out on off over under again further then once here there when where why how all any both each few more most other some such no nor not only own same so than too very can will just should now is am are was were be been being have has had do does did this that these those it its i me my we our you your he she they them their what which who whom as until because also may might would could must shall".split()
)

SENTENCE_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'“‘(])")
WORD_RE = re.compile(r"[\w]+(?:['’][\w]+)?", re.UNICODE)


class _TextExtractor(HTMLParser):
    """Convert HTML content to readable plain text."""

    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.skip_depth = 0

    def handle_starttag(self, tag: str, _attrs: list[tuple[str, str | None]]) -> None:
        del _attrs
        if tag in {"script", "style", "noscript", "svg"}:
            self.skip_depth += 1
        elif tag in {"p", "div", "br", "li", "h1", "h2", "h3", "tr"} and not self.skip_depth:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript", "svg"} and self.skip_depth:
            self.skip_depth -= 1
        elif tag in {"p", "div", "li", "h1", "h2", "h3", "tr"} and not self.skip_depth:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self.skip_depth:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "utf-16", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _clean_text(text: str) -> str:
    text = html.unescape(text).replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\t\f\v ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(filename: str, data: bytes) -> str:
    """Extract readable text from common office, web, image, and text files."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    if suffix in {".txt", ".md", ".markdown", ".rst", ".log", ".yaml", ".yml"}:
        return _clean_text(_decode(data))
    if suffix in {".json", ".xml"}:
        raw = _decode(data)
        try:
            if suffix == ".json":
                return _clean_text(json.dumps(json.loads(raw), indent=2, ensure_ascii=False))
            root = ElementTree.fromstring(raw)
            return _clean_text("\n".join(value for value in root.itertext() if value.strip()))
        except (json.JSONDecodeError, ElementTree.ParseError):
            return _clean_text(raw)
    if suffix in {".html", ".htm"}:
        parser = _TextExtractor()
        parser.feed(_decode(data))
        return _clean_text("".join(parser.parts))
    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        rows = csv.reader(io.StringIO(_decode(data)), delimiter=delimiter)
        return _clean_text("\n".join(" | ".join(cell.strip() for cell in row) for row in rows))
    if suffix == ".rtf":
        try:
            from striprtf.striprtf import rtf_to_text
        except ImportError as exc:
            raise RuntimeError("RTF support needs the optional striprtf package.") from exc
        return _clean_text(rtf_to_text(_decode(data)))
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return _clean_text("\n\n".join(page.extract_text() or "" for page in reader.pages))
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            parts.extend(" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        return _clean_text("\n".join(parts))
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        slides = []
        for index, slide in enumerate(presentation.slides, start=1):
            items = [shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False) and shape.text.strip()]
            slides.append(f"Slide {index}\n" + "\n".join(items))
        return _clean_text("\n\n".join(slides))
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(" | ".join(values))
        return _clean_text("\n".join(parts))
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}:
        try:
            from PIL import Image
            import pytesseract
        except ImportError as exc:
            raise RuntimeError("Image OCR needs Pillow and pytesseract. Install the project requirements.") from exc
        try:
            with Image.open(io.BytesIO(data)) as image:
                return _clean_text(pytesseract.image_to_string(image))
        except pytesseract.TesseractNotFoundError as exc:
            raise RuntimeError("Tesseract OCR is not installed. Install the Tesseract system application to process images.") from exc
    raise ValueError(f"No extractor is available for {suffix}.")


def split_sentences(text: str) -> list[str]:
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    return [part.strip() for part in SENTENCE_SPLIT.split(normalized) if part.strip()]


def analyze_text(text: str) -> dict[str, Any]:
    words = WORD_RE.findall(text)
    sentences = split_sentences(text)
    paragraphs = [part for part in re.split(r"\n\s*\n", text.strip()) if part.strip()]
    tokens = [word.lower() for word in words if len(word) > 2 and word.lower() not in STOP_WORDS]
    frequencies = Counter(tokens)
    syllable_count = sum(_syllables(word) for word in words)
    sentence_count = max(len(sentences), 1)
    word_count = len(words)
    if word_count:
        flesch = 206.835 - 1.015 * (word_count / sentence_count) - 84.6 * (syllable_count / word_count)
    else:
        flesch = 0.0
    return {
        "words": word_count,
        "characters": len(text),
        "characters_no_spaces": len(re.sub(r"\s", "", text)),
        "sentences": len(sentences),
        "paragraphs": len(paragraphs) if text.strip() else 0,
        "reading_minutes": max(1, math.ceil(word_count / 225)) if word_count else 0,
        "speaking_minutes": max(1, math.ceil(word_count / 130)) if word_count else 0,
        "average_word_length": round(sum(map(len, words)) / word_count, 1) if word_count else 0,
        "average_sentence_length": round(word_count / len(sentences), 1) if sentences else 0,
        "flesch_score": round(max(0.0, min(100.0, flesch)), 1) if word_count else 0,
        "keywords": frequencies.most_common(12),
        "sentences_list": sentences,
    }


def _syllables(word: str) -> int:
    cleaned = re.sub(r"[^a-z]", "", word.lower())
    groups = re.findall(r"[aeiouy]+", cleaned)
    count = len(groups)
    if cleaned.endswith("e") and count > 1 and not cleaned.endswith(("le", "ye")):
        count -= 1
    return max(1, count) if cleaned else 0


def summarize_text(
    text: str,
    *,
    style: str = "Balanced",
    length: str = "Standard",
    focus: str = "",
) -> dict[str, Any]:
    """Create an explainable extractive summary, preserving original sentences."""
    analysis = analyze_text(text)
    sentences: list[str] = analysis["sentences_list"]
    if not sentences:
        return {"summary": "", "key_points": [], "source_sentences": 0, "summary_sentences": 0, "compression_percent": 0}

    length_fractions = {"Brief": 0.12, "Standard": 0.22, "Detailed": 0.35}
    target = max(1, min(len(sentences), math.ceil(len(sentences) * length_fractions.get(length, 0.22))))
    frequencies = Counter(
        token.lower()
        for token in WORD_RE.findall(text)
        if len(token) > 2 and token.lower() not in STOP_WORDS
    )
    max_frequency = max(frequencies.values(), default=1)
    focus_terms = {term.lower() for term in WORD_RE.findall(focus) if len(term) > 2}
    action_terms = {"must", "need", "should", "recommend", "recommendation", "action", "next", "deadline", "owner", "propose", "plan", "will"}
    scored: list[tuple[float, int, str]] = []
    for index, sentence in enumerate(sentences):
        tokens = [token.lower() for token in WORD_RE.findall(sentence) if token.lower() in frequencies]
        if not tokens:
            continue
        term_score = sum(frequencies[token] / max_frequency for token in tokens) / math.sqrt(len(tokens))
        position_score = 1.18 if index == 0 else 1.08 if index < max(2, len(sentences) * 0.15) else 1.0
        focus_score = 1 + 0.45 * sum(token in focus_terms for token in tokens) / max(len(focus_terms), 1) if focus_terms else 1.0
        action_score = 1.25 if style == "Action items" and any(token in action_terms for token in tokens) else 1.0
        length_penalty = 0.7 if len(tokens) < 5 else 1.0
        scored.append((term_score * position_score * focus_score * action_score * length_penalty, index, sentence))

    selected = sorted(scored, key=lambda item: item[0], reverse=True)[:target]
    selected.sort(key=lambda item: item[1])
    chosen = [sentence for _, _, sentence in selected] or [sentences[0]]

    style_key = style.lower()
    if style_key == "bullet points":
        summary = "\n".join(f"• {sentence}" for sentence in chosen)
    elif style_key == "executive brief":
        summary = "\n\n".join(("Executive brief", " ".join(chosen)))
    elif style_key == "action items":
        summary = "\n".join(f"☐ {sentence}" for sentence in chosen)
    else:
        summary = " ".join(chosen)

    compression = round((1 - len(WORD_RE.findall(summary)) / max(analysis["words"], 1)) * 100)
    return {
        "summary": summary,
        "key_points": chosen[: min(5, len(chosen))],
        "source_sentences": len(sentences),
        "summary_sentences": len(chosen),
        "compression_percent": max(0, compression),
    }


@dataclass
class DocumentRecord:
    name: str
    text: str
    summary: str
    word_count: int
    file_type: str

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)